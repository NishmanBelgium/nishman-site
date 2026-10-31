#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Retire — ou remet — des références de l'affichage du site.

Pourquoi un drapeau et non une suppression. Une référence effacée des
données emporte avec elle son EAN, son prix, ses cinq traductions et
l'adresse que Google connaît déjà. La remettre devient un chantier, et
son ancienne adresse tombe en erreur. Un drapeau « retire » fait la même
chose aux yeux du visiteur et se défait en une seconde.

Ce que le drapeau retire, et ce qu'il ne retire pas. Il retire la
référence de ce qui s'AFFICHE : le catalogue professionnel, le
formulaire de devis, la boutique, les pages de famille, le plan du site
et le flux Google. Il ne la retire PAS du module de tarification : un
panier déjà constitué, un devis déjà envoyé ou une commande en cours
doivent continuer à calculer leur montant. Un prix qui disparaît en
cours de route, c'est un client qui voit un total et en paie un autre.

Usage :
    python3 outils/retirer.py liste.xlsx        applique le fichier rempli
    python3 outils/retirer.py --etat            liste ce qui est retiré
    python3 outils/retirer.py --remettre EAN    remet une référence

Le fichier attendu est celui envoyé à la cliente : deux onglets, une
colonne « Retirer ? » où OUI signifie retirer. Une référence absente de
la colonne est remise en ligne — le fichier fait foi, en entier.
"""
import io
import json
import sys

RACINE = "/home/claude/nishman-site"
SOURCES = {
    "Catalogue pro (salons)": RACINE + "/assets/data/products.json",
    "Boutique (particuliers)": RACINE + "/assets/data/boutique.json",
}


def charger(chemin):
    d = json.load(io.open(chemin, encoding="utf-8"))
    return d if isinstance(d, list) else (d.get("products") or d.get("produits") or [])


def ecrire(chemin, liste):
    with io.open(chemin, "w", encoding="utf-8") as f:
        json.dump(liste, f, ensure_ascii=False, indent=1)
        f.write("\n")


def etat():
    for nom, chemin in SOURCES.items():
        liste = charger(chemin)
        hors = [p for p in liste if p.get("retire")]
        print("%s : %d retirée(s) sur %d" % (nom, len(hors), len(liste)))
        for p in hors:
            print("   ", p["ean"], "|", p.get("name") or p.get("nom_court") or p.get("nom"))


def appliquer(fichier):
    from openpyxl import load_workbook
    wb = load_workbook(fichier, data_only=True)
    total = 0
    for nom, chemin in SOURCES.items():
        if nom not in wb.sheetnames:
            print("onglet absent, ignoré :", nom)
            continue
        ws = wb[nom]
        entetes = [str(c.value or "").strip() for c in ws[1]]
        try:
            col = entetes.index("Retirer ?")
        except ValueError:
            print("colonne « Retirer ? » absente de", nom, "— ignoré")
            continue
        voulus = set()
        for r in ws.iter_rows(min_row=2):
            v = r[col].value
            if v is not None and str(v).strip().upper().startswith("O"):
                voulus.add(str(r[0].value).strip())

        liste = charger(chemin)
        inconnus = voulus - {str(p["ean"]) for p in liste}
        if inconnus:
            print("EAN inconnus dans", nom, ":", ", ".join(sorted(inconnus)))

        change = 0
        for p in liste:
            avant = bool(p.get("retire"))
            apres = str(p["ean"]) in voulus
            if avant == apres:
                continue
            if apres:
                p["retire"] = True
            else:
                p.pop("retire", None)
            change += 1
        if change:
            ecrire(chemin, liste)
        hors = sum(1 for p in liste if p.get("retire"))
        total += hors
        print("%s : %d retirée(s) sur %d  (%d changement(s))" % (nom, hors, len(liste), change))
    print("total retiré :", total)


def remettre(ean):
    for nom, chemin in SOURCES.items():
        liste = charger(chemin)
        touche = False
        for p in liste:
            if str(p["ean"]) == str(ean) and p.pop("retire", None):
                touche = True
        if touche:
            ecrire(chemin, liste)
            print("remise en ligne dans", nom, ":", ean)


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a or a[0] == "--etat":
        etat()
    elif a[0] == "--remettre" and len(a) > 1:
        remettre(a[1])
    else:
        appliquer(a[0])
